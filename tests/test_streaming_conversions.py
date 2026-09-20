"""Data preservation and bounded-memory regressions for large inputs."""
import csv
import json
import tracemalloc

import pytest
from openpyxl import load_workbook

from cove_converter.engines import data, spreadsheets
from cove_converter.engines.text_io import open_text


def test_oversize_cell_fails_without_replacing_existing_output(tmp_path):
    source, dest = tmp_path / "long.csv", tmp_path / "long.xlsx"
    source.write_text("header\n" + "x" * 40000 + "\n", encoding="utf-8")
    dest.write_bytes(b"existing workbook")
    worker = spreadsheets.SpreadsheetWorker(source, dest)
    errors = []
    worker.failed.connect(lambda message, trace: errors.append(message))
    worker.run()
    assert len(errors) == 1
    assert "Row 2, column 1" in errors[0] and "32767" in errors[0]
    assert dest.read_bytes() == b"existing workbook"
    assert not list(tmp_path.glob(".long.cove-part-*"))


@pytest.mark.parametrize("limit, content, match", [
    ("MAX_XLSX_ROWS", "a\nb\nc\n", "Row 3"),
    ("MAX_XLSX_COLUMNS", "a,b,c\n", "column 3"),
])
def test_sheet_dimensions_are_validated(tmp_path, monkeypatch, limit, content, match):
    monkeypatch.setattr(spreadsheets, limit, 2)
    source = tmp_path / "in.csv"
    source.write_text(content, encoding="utf-8")
    with pytest.raises(spreadsheets.SpreadsheetLimitError, match=match):
        spreadsheets._csv_to_xlsx(source, tmp_path / "out.xlsx")


def test_column_limit_accepts_exact_boundary(tmp_path, monkeypatch):
    # The overflow message hardcodes MAX_XLSX_COLUMNS + 1 regardless of the
    # actual row length, so a row of 3 columns against a limit of 2 can't by
    # itself distinguish a correct `>` check from an off-by-one `>=` bug -
    # both raise the same "column 3" message. A row at exactly the limit
    # must be accepted to prove the boundary is inclusive.
    monkeypatch.setattr(spreadsheets, "MAX_XLSX_COLUMNS", 2)
    source = tmp_path / "in.csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")
    dest = tmp_path / "out.xlsx"
    spreadsheets._csv_to_xlsx(source, dest)
    wb = load_workbook(dest, read_only=True)
    try:
        assert [c.value for c in next(wb.active.iter_rows())] == ["a", "b"]
    finally:
        wb.close()


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32", "cp1252", "latin-1"])
def test_streamed_csv_preserves_encoding_multiline_and_formulas(tmp_path, encoding):
    source, dest = tmp_path / "in.csv", tmp_path / "out.xlsx"
    text = "\x81" if encoding == "latin-1" else "Café"
    with source.open("w", encoding=encoding, newline="") as handle:
        csv.writer(handle).writerow([text, "first\r\nsecond", "=1+1", "x" * 32767])
    spreadsheets._csv_to_xlsx(source, dest)
    wb = load_workbook(dest, read_only=True)
    try:
        cells = next(wb.active.iter_rows())
        assert [c.value for c in cells] == [text, "first\r\nsecond", "=1+1", "x" * 32767]
        assert cells[2].data_type == "s"
    finally:
        wb.close()


def test_encoding_fallback_handles_error_after_first_chunk(tmp_path):
    path = tmp_path / "late.csv"
    path.write_bytes(b"x" * 70000 + b"\x93quoted\x94")
    with open_text(path, legacy_fallback=True) as source:
        assert source.read().endswith("“quoted”")


@pytest.mark.parametrize("suffix", [".json", ".jsonl"])
def test_ndjson_conversion_has_bounded_memory(tmp_path, suffix):
    source, dest = tmp_path / "records.ndjson", tmp_path / ("out" + suffix)
    with source.open("w", encoding="utf-8") as output:
        for index in range(4000):
            output.write(json.dumps({"index": index, "text": "x" * 1024}) + "\n")
    worker = data.DataWorker(source, dest)
    tracemalloc.start()
    try:
        worker._convert()
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < 2 * 1024 * 1024
    if suffix == ".json":
        records = json.loads(dest.read_text())
    else:
        records = [json.loads(line) for line in dest.read_text().splitlines()]
    assert len(records) == 4000 and records[-1]["index"] == 3999


def test_late_ndjson_error_preserves_destination(tmp_path):
    source, dest = tmp_path / "in.jsonl", tmp_path / "out.json"
    source.write_text('{"ok": 1}\n\n{"x":1,"x":2}\n')
    dest.write_text("keep me")
    worker = data.DataWorker(source, dest)
    errors = []
    worker.failed.connect(lambda message, trace: errors.append(message))
    worker.run()
    assert "line 3" in errors[0]
    assert dest.read_text() == "keep me"


def test_yaml_duplicate_key_check_is_not_quadratic():
    """Deterministic operation-count oracle for the YAML duplicate-key scan.

    A per-key linear scan against every previously seen key (``for prev in
    seen_python: prev == key``) performs n(n-1)/2 equality comparisons for n
    distinct keys. Set/dict membership performs O(1) comparisons per key via
    hashing. Counting real ``__eq__`` calls avoids a noisy wall-clock timing
    oracle while still failing against the former quadratic implementation.
    """
    eq_calls = [0]

    class _CountedKey:
        def __init__(self, index):
            self.index = index

        def __hash__(self):
            return hash(self.index)

        def __eq__(self, other):
            eq_calls[0] += 1
            return isinstance(other, _CountedKey) and self.index == other.index

    loader_cls = data._build_collision_loader()

    def _construct_counted(loader, node):
        return _CountedKey(int(loader.construct_scalar(node)))

    loader_cls.add_constructor("!counted", _construct_counted)

    key_count = 400
    text = "\n".join(f"? !counted {i}\n: {i}" for i in range(key_count)) + "\n"
    loader = loader_cls(text)
    try:
        result = loader.get_single_data()
    finally:
        loader.dispose()

    assert len(result) == key_count
    # O(n) membership tracking stays far below this; O(n^2) scanning blows
    # past it (400 keys: linear ~0, quadratic ~79,800 comparisons).
    assert eq_calls[0] < key_count * 4, (
        f"{eq_calls[0]} equality comparisons for {key_count} distinct keys "
        "indicates quadratic duplicate-key scanning"
    )


def test_write_only_workbook_has_bounded_memory(tmp_path):
    # 20,000 rows keeps the input reader's per-row footprint sensitive: a true
    # incremental reader stays flat regardless of file size (~0.8 MiB
    # measured), while buffering the whole decoded CSV text in memory before
    # parsing scales with input size (~10.8 MiB measured at this row count).
    # The 3 MiB threshold sits with wide margin on both sides of that gap.
    row_count = 20000
    source, dest = tmp_path / "many.csv", tmp_path / "many.xlsx"
    with source.open("w", encoding="utf-8") as output:
        for index in range(row_count):
            output.write(",".join(f"row{index}-col{col}" for col in range(8)) + "\n")
    tracemalloc.start()
    try:
        spreadsheets._csv_to_xlsx(source, dest)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < 3 * 1024 * 1024
    wb = load_workbook(dest, read_only=True)
    try:
        assert sum(1 for _ in wb.active.rows) == row_count
    finally:
        wb.close()
